"""Unit tests for the Enhancv-inspired ATS Compatibility and Resume Checker system."""

from __future__ import annotations

from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document

from app.agents.ats_agent import ATSAgent
from app.agents.models import AnalysisState, FormattingRisk
from app.agents.skill_agent import SkillAgent
from app.agents.evidence_agent import EvidenceAgent
from app.parsing.documents import extract_document_text_with_diagnostics
from app.parsing.models import (
    ContactInformation,
    ExtractedItem,
    JobDescriptionExtraction,
    ResumeExtraction,
    ResumeLayoutDiagnostics,
)
from app.parsing.resume import parse_resume
from app.parsing.validation import validate_upload
from app.scoring import ScoringEngine


def create_sample_docx(has_table: bool = True) -> bytes:
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("jane.doe@example.com | +1 555 123 4567 | San Francisco, CA")
    doc.add_paragraph("https://linkedin.com/in/janedoe | https://github.com/janedoe")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph("Full Stack Software Engineer with 5+ years building web applications.")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, TypeScript, React, PostgreSQL, Docker, AWS")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Software Engineer at TechCorp (2020 - Present)")
    doc.add_paragraph("Responsible for managing backend services and assisted in database setup.")
    doc.add_paragraph("Engineered real-time telemetry pipeline scaling throughput by 40% and saving $25K annually.")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("BS in Computer Science, Stanford University (2016 - 2020)")

    if has_table:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Skill"
        table.cell(0, 1).text = "Proficiency"
        table.cell(1, 0).text = "Python"
        table.cell(1, 1).text = "Expert"

    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def minimal_pdf_bytes(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    body = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    body += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    body += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return body


def test_docx_diagnostics_detection() -> None:
    docx_bytes = create_sample_docx(has_table=True)
    upload = validate_upload("resume.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_bytes, 500_000)
    text, diagnostics = extract_document_text_with_diagnostics(upload)

    assert "Jane Doe" in text
    assert diagnostics.has_tables is True
    assert diagnostics.table_count >= 1
    assert any("table" in warning.lower() for warning in diagnostics.formatting_warnings)


def test_pdf_diagnostics_detection() -> None:
    pdf_bytes = minimal_pdf_bytes("Jane Doe Engineer")
    upload = validate_upload("resume.pdf", "application/pdf", pdf_bytes, 500_000)
    text, diagnostics = extract_document_text_with_diagnostics(upload)

    assert "Jane Doe Engineer" in text
    assert diagnostics.page_count == 1
    assert diagnostics.is_scanned is False


def test_standalone_ats_agent_five_pillars() -> None:
    """Verify that ATSAgent conducts a comprehensive 5-pillar audit without requiring a job description."""
    owner_id = uuid4()
    docx_bytes = create_sample_docx(has_table=False)
    upload = validate_upload("resume.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_bytes, 500_000)
    text, diagnostics = extract_document_text_with_diagnostics(upload)
    extraction = parse_resume(text, "resume.docx", diagnostics=diagnostics)

    general_jd = JobDescriptionExtraction(
        title="General Professional Benchmark",
        requirements=[],
        source="system_benchmark",
        raw_text="General Professional ATS Resume Quality Benchmark",
    )

    state = AnalysisState(
        analysis_id=uuid4(),
        owner_id=owner_id,
        resume_extraction=extraction,
        job_extraction=general_jd,
        evidence_corpus=extraction.items,
        grounded_contexts=[],
    )

    ats_agent = ATSAgent()
    output = ats_agent.analyze(state)

    # Validate output structure
    assert output.overall_ats_score >= 0.0
    assert output.readability_score >= 0.0
    assert output.content_quality_score >= 0.0
    assert output.structure_score >= 0.0

    # Check 5 pillars are present in category_pillars (standalone mode has 5 general pillars)
    pillars = output.category_pillars
    assert "readability" in pillars
    assert "structure" in pillars
    assert "content" in pillars
    assert "contact" in pillars
    assert "skills" in pillars

    # Readability pillar
    assert pillars["readability"]["status"] in ("pass", "warn", "fail")
    assert "score" in pillars["readability"]

    # Structure pillar
    assert pillars["structure"]["status"] in ("pass", "warn", "fail")
    assert "issues_count" in pillars["structure"]

    # Content quality pillar
    assert pillars["content"]["status"] in ("pass", "warn", "fail")
    assert output.weak_verbs_found is not None

    # Contact info pillar
    assert pillars["contact"]["status"] in ("pass", "warn", "fail")

    # Prioritized recommendations
    assert len(output.recommendations) > 0
    # Check that recommendation priorities exist
    priorities = {rec.priority.lower() for rec in output.recommendations}
    assert priorities.intersection({"critical", "high", "medium", "recommended", "optional"})


def test_weak_verb_and_quantifiable_metric_detection() -> None:
    """Verify weak passive verbs and metric coverage are correctly identified."""
    owner_id = uuid4()
    items = [
        ExtractedItem(
            id=uuid4(),
            type="experience",
            text="Responsible for handling daily customer tickets and assisted the manager.",
            source="resume.pdf",
            location="Experience > Role 1 > Bullet 1",
            confidence=0.9,
        ),
        ExtractedItem(
            id=uuid4(),
            type="experience",
            text="Increased page load speed by 35% resulting in $120,000 extra conversion revenue.",
            source="resume.pdf",
            location="Experience > Role 1 > Bullet 2",
            confidence=0.95,
        ),
    ]

    extraction = ResumeExtraction(
        source="resume.pdf",
        name="Alex Smith",
        contact_information=ContactInformation(
            email="alex@example.com",
            phone="+1 555 333 4444",
            location="New York, NY",
            links=["https://linkedin.com/in/alexsmith"],
        ),
        summary="Experienced engineer with strong background.",
        sections={
            "summary": [ExtractedItem(id=uuid4(), type="summary", text="Experienced engineer", source="resume.pdf", location="Summary", confidence=0.9)],
            "experience": items,
            "skills": [ExtractedItem(id=uuid4(), type="skill", text="Python, SQL", source="resume.pdf", location="Skills", confidence=0.9)],
            "education": [ExtractedItem(id=uuid4(), type="education", text="BS CS", source="resume.pdf", location="Education", confidence=0.9)],
        },
        items=items,
        raw_text="Alex Smith\nalex@example.com\n\nExperience\n- Responsible for handling daily customer tickets and assisted the manager.\n- Increased page load speed by 35% resulting in $120,000 extra conversion revenue.",
    )

    state = AnalysisState(
        analysis_id=uuid4(),
        owner_id=owner_id,
        resume_extraction=extraction,
        job_extraction=JobDescriptionExtraction(title="General Benchmark", requirements=[], source="bench", raw_text="General Benchmark Description"),
        evidence_corpus=items,
        grounded_contexts=[],
    )

    output = ATSAgent().analyze(state)

    # Weak verbs should detect "responsible for" and "assisted"
    assert len(output.weak_verbs_found) >= 2
    # Quantified ratio should detect metrics in bullets
    assert 0.2 <= output.quantified_bullets_ratio <= 1.0

    # Action verb recommendation should exist
    verb_recs = [r for r in output.recommendations if "verb" in r.suggested_action.lower() or "verb" in r.title.lower()]
    assert len(verb_recs) >= 1


from collections.abc import AsyncGenerator
import httpx
import pytest_asyncio
from redis.asyncio import Redis
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.core.config import get_settings
from app.db import get_db
from app.main import app
from app.services.redis import get_redis

settings = get_settings()


def database_url() -> str:
    return settings.database_url.replace("postgresql://", "postgresql+asyncpg://", 1)


@pytest_asyncio.fixture
async def integration_resources() -> AsyncGenerator[tuple[async_sessionmaker[AsyncSession], Redis], None]:
    engine = create_async_engine(database_url(), pool_pre_ping=True)
    redis = Redis.from_url(settings.redis_url, decode_responses=True)
    try:
        async with engine.connect() as connection:
            await connection.execute(text("SELECT 1"))
        await redis.ping()
    except Exception as error:
        await redis.aclose()
        await engine.dispose()
        pytest.skip(f"Integration tests require PostgreSQL and Redis: {error}")

    yield async_sessionmaker(engine, expire_on_commit=False), redis
    await redis.flushdb()
    await redis.aclose()
    await engine.dispose()


@pytest_asyncio.fixture
async def client(integration_resources: tuple[async_sessionmaker[AsyncSession], Redis]) -> AsyncGenerator[httpx.AsyncClient, None]:
    session_factory, redis = integration_resources

    async def override_db() -> AsyncGenerator[AsyncSession, None]:
        async with session_factory() as session:
            yield session

    async def override_redis() -> AsyncGenerator[Redis, None]:
        yield redis

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_redis] = override_redis
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_standalone_ats_api_workflow_and_delete(client: httpx.AsyncClient) -> None:
    """Verify standalone ATS check without JD and subsequent analysis deletion via REST API."""
    user_email = f"ats_tester_{uuid4()}@example.com"
    reg = await client.post(
        "/api/auth/signup",
        json={"email": user_email, "password": "TestPassword123!", "full_name": "ATS Candidate"},
    )
    assert reg.status_code == 201
    token = reg.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    docx_bytes = create_sample_docx(has_table=False)
    files = {"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    upload_resp = await client.post("/api/ingestion/resumes", files=files, headers=headers)
    assert upload_resp.status_code == 201
    resume_id = upload_resp.json()["id"]

    # Trigger standalone ATS analysis without job_description_id
    baseline_resp = await client.post(
        "/api/analysis/baseline",
        json={"resume_id": resume_id},
        headers=headers,
    )
    assert baseline_resp.status_code == 201
    report = baseline_resp.json()
    analysis_id = report["analysis_id"]

    assert report["has_target_job"] is False
    assert "readability" in report["category_pillars"]
    assert "content" in report["category_pillars"]
    assert len(report["recommendations"]) > 0

    # Verify report is retrievable
    get_resp = await client.get(f"/api/analysis/{analysis_id}/baseline", headers=headers)
    assert get_resp.status_code == 200

    # Delete the analysis
    del_resp = await client.delete(f"/api/analysis/{analysis_id}", headers=headers)
    assert del_resp.status_code == 204

    # Verify analysis is gone
    get_after_del = await client.get(f"/api/analysis/{analysis_id}/baseline", headers=headers)
    assert get_after_del.status_code == 404

