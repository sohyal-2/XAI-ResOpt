from io import BytesIO

import pytest
from docx import Document
from pydantic import ValidationError

from app.parsing.documents import extract_document_text
from app.parsing.job_description import parse_job_description
from app.parsing.models import ExtractedItem
from app.parsing.resume import parse_resume
from app.parsing.url_safety import UnsafeUrlError, validate_public_url
from app.parsing.validation import CorruptDocumentError, UnsupportedFileError, validate_upload


RESUME_TEXT = """Ada Lovelace
ada@example.com | +1 555 123 4567
https://github.com/ada

Summary
Backend engineer building reliable APIs.

Skills
Python, FastAPI, PostgreSQL

Experience
- Built APIs and improved service reliability.

Projects
- Analytical Engine: Python project.

Education
Babbage University, Computer Science

Certifications
Cloud Certificate

Achievements
Engineering award

Publications
Reliable Systems paper
"""


def minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    body = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    body += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    body += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return body


def test_plain_text_resume_extracts_sections_and_provenance() -> None:
    upload = validate_upload("resume.txt", "text/plain", RESUME_TEXT.encode(), 100_000)
    extraction = parse_resume(extract_document_text(upload), upload.filename)
    assert extraction.name == "Ada Lovelace"
    assert extraction.contact_information.email == "ada@example.com"
    assert {"skills", "experience", "projects", "education", "certifications", "achievements", "publications"} <= set(extraction.sections)
    assert all(item.source == "resume.txt" and item.location and 0 <= item.confidence <= 1 for item in extraction.items)


def test_pdf_resume_is_readable() -> None:
    upload = validate_upload("resume.pdf", "application/pdf", minimal_pdf("Ada Lovelace"), 100_000)
    assert "Ada Lovelace" in extract_document_text(upload)


def test_docx_resume_is_readable() -> None:
    document = Document()
    document.add_paragraph("Ada Lovelace")
    document.add_paragraph("Skills")
    document.add_paragraph("Python")
    stream = BytesIO()
    document.save(stream)
    upload = validate_upload("resume.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", stream.getvalue(), 100_000)
    assert "Ada Lovelace" in extract_document_text(upload)


def test_job_description_text_extracts_requirements() -> None:
    extraction = parse_job_description("Senior Python Engineer\nRequirements:\n- Python and FastAPI experience\n- PostgreSQL knowledge", "pasted job description")
    assert extraction.title == "Senior Python Engineer"
    assert len(extraction.requirements) == 2
    assert "python" in extraction.requirements[0].normalized_terms


def test_invalid_and_corrupt_files_are_rejected() -> None:
    with pytest.raises(UnsupportedFileError):
        validate_upload("resume.exe", "application/octet-stream", b"bad", 100_000)
    with pytest.raises(CorruptDocumentError):
        validate_upload("resume.pdf", "application/pdf", b"not a pdf", 100_000)
    with pytest.raises(CorruptDocumentError):
        validate_upload("resume.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"not a docx", 100_000)


def test_invalid_urls_and_ssrf_targets_are_rejected() -> None:
    for url in ("ftp://example.com/file", "http://127.0.0.1:8000", "http://localhost", "http://10.0.0.1"):
        with pytest.raises(UnsafeUrlError):
            validate_public_url(url)


def test_extraction_schema_rejects_invalid_confidence() -> None:
    with pytest.raises(ValidationError):
        ExtractedItem(type="skill", text="Python", source="resume.txt", location="Skills > line 1", confidence=1.5)
