"""Tests for deterministic ATS resume document builder, pre-rendering validator, and PDF/DOCX exporters."""

from collections.abc import AsyncGenerator
import io
from uuid import uuid4

import httpx
from pypdf import PdfReader
from docx import Document
import pytest
import pytest_asyncio
from redis.asyncio import Redis
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.core.config import get_settings
from app.db import get_db
from app.main import app
from app.models.entities import ChangeDecision, ResumeChange
from app.schemas.rendering import (
    DeterministicResumeDocument,
    ResumeContactBlock,
    ResumeSectionBlock,
)
from app.services.redis import get_redis
from app.services.rendering import (
    ATSDocxRenderer,
    ATSPdfRenderer,
    build_deterministic_document,
    validate_resume_for_rendering,
)

settings = get_settings()


def database_url() -> str:
    return settings.database_url.replace("postgresql://", "postgresql+asyncpg://", 1)


@pytest_asyncio.fixture
async def integration_resources() -> AsyncGenerator[tuple[async_sessionmaker[AsyncSession], Redis], None]:
    engine = create_async_engine(database_url(), pool_pre_ping=True)
    redis = Redis.from_url(settings.redis_url, decode_responses=True)
    try:
        async with engine.connect() as connection:
            await connection.execute(text("SELECT 1"))
        await redis.ping()
    except Exception as error:
        await redis.aclose()
        await engine.dispose()
        pytest.skip(f"Resume rendering integration tests require PostgreSQL and Redis: {error}")

    yield async_sessionmaker(engine, expire_on_commit=False), redis
    await redis.flushdb()
    await redis.aclose()
    await engine.dispose()


@pytest_asyncio.fixture
async def client(integration_resources: tuple[async_sessionmaker[AsyncSession], Redis]) -> AsyncGenerator[httpx.AsyncClient, None]:
    session_factory, redis = integration_resources

    async def override_db() -> AsyncGenerator[AsyncSession, None]:
        async with session_factory() as session:
            yield session

    async def override_redis() -> AsyncGenerator[Redis, None]:
        yield redis

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_redis] = override_redis
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as test_client:
        yield test_client
    app.dependency_overrides.clear()


SAMPLE_STRUCTURED_TEXT = """Jordan Taylor
jordan.taylor@example.com | +1 555 987 6543 | Austin, TX | https://github.com/jordantaylor | https://linkedin.com/in/jordantaylor

Professional Summary
Senior Software Engineer with 7+ years of experience architecting distributed backend microservices and databases.

Work Experience
Senior Backend Engineer at CloudScale Inc
- Architected high-throughput REST APIs using Python, FastAPI, and PostgreSQL.
- Optimized PostgreSQL database performance, reducing p99 latency by 35%.
- Implemented automated CI/CD pipelines using Docker.

Technical Skills
Python, FastAPI, PostgreSQL, Docker, Redis, Kubernetes, SQL, Git

Education
University of Texas at Austin - BS in Computer Science
"""

SAMPLE_JOB_TEXT = """Senior Backend Engineer
Requirements
- 5+ years experience with Python and FastAPI
- Production experience with PostgreSQL and Docker
"""


def test_deterministic_document_builder() -> None:
    """Verifies that builder deterministically extracts headers, sections, bullets, and metadata."""
    doc = build_deterministic_document(SAMPLE_STRUCTURED_TEXT)

    # 1. Verify Header
    assert doc.header.name == "Jordan Taylor"
    assert doc.header.email == "jordan.taylor@example.com"
    assert doc.header.phone == "+1 555 987 6543"
    assert "Austin, TX" in (doc.header.location or "")
    assert "github.com" in (doc.header.github or "")
    assert "linkedin.com" in (doc.header.linkedin or "")

    # 2. Verify Sections & Normalization
    section_keys = [s.section_key for s in doc.sections]
    assert "summary" in section_keys
    assert "experience" in section_keys
    assert "skills" in section_keys
    assert "education" in section_keys

    # Check bullets in experience
    exp_sec = next(s for s in doc.sections if s.section_key == "experience")
    all_bullets = [b for item in exp_sec.items for b in item.bullets]
    assert len(all_bullets) == 3
    assert "Architected high-throughput REST APIs" in all_bullets[0]


def test_pre_rendering_validation_rules() -> None:
    """Tests all 5 mandatory safety checks in the pre-rendering validation layer."""
    doc = build_deterministic_document(SAMPLE_STRUCTURED_TEXT)

    # 1. Valid resume passes
    val = validate_resume_for_rendering(doc, SAMPLE_STRUCTURED_TEXT)
    assert val.is_valid is True
    assert all(val.checks_passed.values())
    assert len(val.issues) == 0

    # 2. Catch missing contact info
    bad_contact_doc = DeterministicResumeDocument(
        header=ResumeContactBlock(name="", email="not-an-email"),
        sections=doc.sections,
    )
    val_bad_contact = validate_resume_for_rendering(bad_contact_doc, SAMPLE_STRUCTURED_TEXT)
    assert val_bad_contact.is_valid is False
    assert val_bad_contact.checks_passed["contact_info_intact"] is False

    # 3. Catch empty accidental sections
    empty_sec_doc = DeterministicResumeDocument(
        header=doc.header,
        sections=[ResumeSectionBlock(section_key="empty_sec", heading="Empty Section", items=[])],
    )
    val_empty = validate_resume_for_rendering(empty_sec_doc, SAMPLE_STRUCTURED_TEXT)
    assert val_empty.is_valid is False
    assert val_empty.checks_passed["no_empty_sections"] is False

    # 4. Catch missing approved changes
    approved_change = ResumeChange(
        id=uuid4(),
        optimization_run_id=uuid4(),
        resume_version_id=uuid4(),
        before_text="Old bullet",
        after_text="Nonexistent approved rewrite text that must be in final resume",
        decision=ChangeDecision.accepted,
        supported=True,
    )
    val_missing_approved = validate_resume_for_rendering(
        doc, SAMPLE_STRUCTURED_TEXT, changes=[approved_change]
    )
    assert val_missing_approved.is_valid is False
    assert val_missing_approved.checks_passed["approved_changes_present"] is False

    # 5. Catch present rejected changes
    rejected_change = ResumeChange(
        id=uuid4(),
        optimization_run_id=uuid4(),
        resume_version_id=uuid4(),
        before_text="Old bullet",
        after_text="Architected high-throughput REST APIs",  # Present in sample text!
        decision=ChangeDecision.rejected,
        supported=False,
    )
    val_rejected_present = validate_resume_for_rendering(
        doc, SAMPLE_STRUCTURED_TEXT, changes=[rejected_change]
    )
    assert val_rejected_present.is_valid is False
    assert val_rejected_present.checks_passed["rejected_changes_absent"] is False


def test_pdf_renderer_produces_valid_pdf() -> None:
    """Tests that ATSPdfRenderer produces valid, parseable single-column PDF binary output."""
    doc = build_deterministic_document(SAMPLE_STRUCTURED_TEXT)
    renderer = ATSPdfRenderer(doc)
    pdf_bytes = renderer.render_pdf_bytes()

    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 1000

    # Parse with PyPDF to verify selectable text content
    reader = PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1
    extracted_text = "\n".join(page.extract_text() or "" for page in reader.pages)

    assert "Jordan Taylor" in extracted_text
    assert "jordan.taylor@example.com" in extracted_text
    assert "PROFESSIONAL SUMMARY" in extracted_text.upper()
    assert "WORK EXPERIENCE" in extracted_text.upper()
    assert "TECHNICAL SKILLS" in extracted_text.upper()


def test_docx_renderer_produces_valid_docx() -> None:
    """Tests that ATSDocxRenderer produces valid, semantic Word document binary output."""
    doc = build_deterministic_document(SAMPLE_STRUCTURED_TEXT)
    renderer = ATSDocxRenderer(doc)
    docx_bytes = renderer.render_docx_bytes()

    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 1000

    # Open with python-docx to verify paragraph hierarchy and styles
    word_doc = Document(io.BytesIO(docx_bytes))
    all_paras = [p.text for p in word_doc.paragraphs]
    combined_doc_text = "\n".join(all_paras)

    assert "Jordan Taylor" in combined_doc_text
    assert "jordan.taylor@example.com" in combined_doc_text
    assert "PROFESSIONAL SUMMARY" in combined_doc_text
    assert "WORK EXPERIENCE" in combined_doc_text
    assert "Architected high-throughput REST APIs using Python, FastAPI, and PostgreSQL." in combined_doc_text


@pytest.mark.asyncio
async def test_resume_export_api_endpoints_and_isolation(client: httpx.AsyncClient) -> None:
    """End-to-end integration test verifying validation, preview, PDF export, DOCX export, and security isolation."""
    # 1. Register & login
    user_email = f"exporter_{uuid4()}@example.com"
    reg_response = await client.post(
        "/api/auth/signup",
        json={"email": user_email, "password": "Password123!", "full_name": "Export Tester"},
    )
    assert reg_response.status_code == 201
    auth_token = reg_response.json()["access_token"]
    headers = {"Authorization": f"Bearer {auth_token}"}

    # 2. Ingest Resume & Job Description
    files = {"file": ("jordan_resume.txt", SAMPLE_STRUCTURED_TEXT.encode("utf-8"), "text/plain")}
    resume_resp = await client.post("/api/ingestion/resumes", files=files, headers=headers)
    assert resume_resp.status_code == 201
    resume_id = resume_resp.json()["id"]

    job_resp = await client.post(
        "/api/ingestion/job-descriptions",
        data={"text": SAMPLE_JOB_TEXT, "title": "Senior Backend Engineer"},
        headers=headers,
    )
    assert job_resp.status_code == 201
    job_id = job_resp.json()["id"]

    # 3. Baseline -> Optimize -> Finalize Review
    base_resp = await client.post("/api/analysis/baseline", json={"resume_id": resume_id, "job_description_id": job_id}, headers=headers)
    analysis_id = base_resp.json()["analysis_id"]

    await client.post(f"/api/analysis/{analysis_id}/optimize", json={"max_iterations": 1}, headers=headers)

    fin_resp = await client.post(f"/api/analysis/{analysis_id}/review/finalize", json={"force_unresolved_claims": True}, headers=headers)
    assert fin_resp.status_code == 200
    new_version_id = fin_resp.json()["new_resume_version_id"]

    # 4. Test GET /api/resumes/versions/{version_id}/validate
    val_resp = await client.get(f"/api/resumes/versions/{new_version_id}/validate", headers=headers)
    assert val_resp.status_code == 200
    val_data = val_resp.json()
    assert val_data["is_valid"] is True
    assert val_data["checks_passed"]["sections_present"] is True
    assert val_data["checks_passed"]["no_empty_sections"] is True

    # 5. Test GET /api/resumes/versions/{version_id}/preview
    prev_resp = await client.get(f"/api/resumes/versions/{new_version_id}/preview", headers=headers)
    assert prev_resp.status_code == 200
    prev_data = prev_resp.json()
    assert prev_data["header"]["name"] == "Jordan Taylor"
    assert len(prev_data["sections"]) >= 3

    # 6. Test GET /api/resumes/versions/{version_id}/export/pdf
    pdf_resp = await client.get(f"/api/resumes/versions/{new_version_id}/export/pdf", headers=headers)
    assert pdf_resp.status_code == 200
    assert pdf_resp.headers["content-type"] == "application/pdf"
    assert len(pdf_resp.content) > 1000
    assert pdf_resp.content.startswith(b"%PDF")

    # 7. Test GET /api/resumes/versions/{version_id}/export/docx
    docx_resp = await client.get(f"/api/resumes/versions/{new_version_id}/export/docx", headers=headers)
    assert docx_resp.status_code == 200
    assert "wordprocessingml" in docx_resp.headers["content-type"]
    assert len(docx_resp.content) > 1000

    # 8. Tenant Isolation Check: Another user cannot export or preview this version
    other_email = f"other_{uuid4()}@example.com"
    other_reg = await client.post("/api/auth/signup", json={"email": other_email, "password": "Password123!", "full_name": "Other"})
    other_headers = {"Authorization": f"Bearer {other_reg.json()['access_token']}"}

    other_pdf = await client.get(f"/api/resumes/versions/{new_version_id}/export/pdf", headers=other_headers)
    assert other_pdf.status_code == 404

    other_docx = await client.get(f"/api/resumes/versions/{new_version_id}/export/docx", headers=other_headers)
    assert other_docx.status_code == 404

    # 9. Direct browser download test via ?token= query parameter (without Authorization header)
    query_pdf = await client.get(f"/api/resumes/versions/{new_version_id}/export/pdf?token={auth_token}")
    assert query_pdf.status_code == 200
    assert query_pdf.headers["content-type"] == "application/pdf"
    assert query_pdf.content.startswith(b"%PDF")

    query_docx = await client.get(f"/api/resumes/versions/{new_version_id}/export/docx?token={auth_token}")
    assert query_docx.status_code == 200
    assert "wordprocessingml" in query_docx.headers["content-type"]
