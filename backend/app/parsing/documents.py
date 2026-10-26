from io import BytesIO
import re

from docx import Document
from pypdf import PdfReader

from app.parsing.models import ResumeLayoutDiagnostics
from app.parsing.validation import CorruptDocumentError, ValidatedUpload

# Security limits against decompression / parser bombs
MAX_PDF_PAGES = 50
MAX_EXTRACTED_CHARACTERS = 250_000


def extract_document_text(upload: ValidatedUpload) -> str:
    text, _ = extract_document_text_with_diagnostics(upload)
    return text


def extract_document_text_with_diagnostics(upload: ValidatedUpload) -> tuple[str, ResumeLayoutDiagnostics]:
    try:
        if upload.filename.lower().endswith(".pdf"):
            text, diag = _extract_pdf_with_diagnostics(upload.content)
        elif upload.filename.lower().endswith(".docx"):
            text, diag = _extract_docx_with_diagnostics(upload.content)
        else:
            text = upload.content.decode("utf-8-sig").strip()
            diag = ResumeLayoutDiagnostics()
            if not text:
                raise CorruptDocumentError("The uploaded file contains no readable text.")

        if len(text) > MAX_EXTRACTED_CHARACTERS:
            text = text[:MAX_EXTRACTED_CHARACTERS]
        return text, diag
    except (OSError, UnicodeDecodeError, ValueError, KeyError) as error:
        raise CorruptDocumentError(f"Unable to read {upload.filename}: {error}") from error


def _extract_pdf_with_diagnostics(content: bytes) -> tuple[str, ResumeLayoutDiagnostics]:
    diag = ResumeLayoutDiagnostics()
    try:
        reader = PdfReader(BytesIO(content), strict=True)
        page_count = len(reader.pages)
        diag.page_count = page_count
        if page_count > MAX_PDF_PAGES:
            raise CorruptDocumentError(f"PDF exceeds maximum supported page limit of {MAX_PDF_PAGES} pages.")
        
        pages_text = []
        total_len = 0
        multi_col_detected = False

        for page in reader.pages:
            x_positions: list[float] = []

            def visitor_body(text_chunk: str, cm: any, tm: any, font_dict: any, font_size: any) -> None:
                if text_chunk and text_chunk.strip():
                    try:
                        x_positions.append(float(tm[4]))
                    except (IndexError, TypeError, ValueError):
                        pass

            try:
                extracted = page.extract_text(visitor_text=visitor_body) or ""
            except Exception:
                extracted = page.extract_text() or ""

            pages_text.append(extracted)
            total_len += len(extracted)

            # Spatial column detection
            if len(x_positions) >= 20:
                left_count = sum(1 for x in x_positions if x < 260)
                right_count = sum(1 for x in x_positions if x > 310)
                if left_count >= 8 and right_count >= 8 and (left_count + right_count) / len(x_positions) > 0.6:
                    multi_col_detected = True

            if total_len > MAX_EXTRACTED_CHARACTERS:
                break

        text = "\n".join(pages_text).strip()
    except CorruptDocumentError:
        raise
    except Exception as error:
        raise CorruptDocumentError("The PDF could not be parsed.") from error

    if not text or not text.strip():
        diag.is_scanned = True
        diag.formatting_warnings.append("Document appears to be a scanned image or lacks extractable digital text.")
        raise CorruptDocumentError("The PDF contains no extractable text. Scanned resumes without a digital text layer cannot be parsed by ATS engines.")

    non_standard = len(re.findall(r"[^\x00-\x7F\u2022\u2013\u2014\u00A0]", text))
    diag.non_standard_char_count = non_standard
    if non_standard > 25:
        diag.formatting_warnings.append(f"Detected {non_standard} non-standard characters or symbols that risk ATS encoding issues.")

    if multi_col_detected:
        diag.has_multi_column = True
        diag.formatting_warnings.append("Multi-column layout detected. ATS parsers typically parse left-to-right across columns, potentially jumbling text.")

    return text, diag


def _extract_docx_with_diagnostics(content: bytes) -> tuple[str, ResumeLayoutDiagnostics]:
    diag = ResumeLayoutDiagnostics()
    try:
        document = Document(BytesIO(content))
        paragraphs = []
        total_len = 0
        for paragraph in document.paragraphs:
            cleaned = paragraph.text.strip()
            if cleaned:
                paragraphs.append(cleaned)
                total_len += len(cleaned)
                if total_len > MAX_EXTRACTED_CHARACTERS:
                    break

        table_text = []
        table_count = len(document.tables)
        diag.table_count = table_count
        if table_count > 0:
            diag.has_tables = True
            diag.formatting_warnings.append(f"Document contains {table_count} table(s). Complex tables can confuse ATS parsers.")

        if total_len < MAX_EXTRACTED_CHARACTERS:
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cleaned = cell.text.strip()
                        if cleaned:
                            table_text.append(cleaned)
                            total_len += len(cleaned)
                            if total_len > MAX_EXTRACTED_CHARACTERS:
                                break
        text = "\n".join(paragraphs + table_text).strip()

        try:
            xml_str = document.element.xml
            if "w:txbx" in xml_str or "v:textbox" in xml_str:
                diag.has_text_boxes = True
                diag.formatting_warnings.append("Document contains text boxes. Standard ATS parsers often skip text box content.")
        except Exception:
            pass

    except Exception as error:
        raise CorruptDocumentError("The DOCX could not be parsed.") from error
    if not text:
        raise CorruptDocumentError("The DOCX contains no extractable text.")

    non_standard = len(re.findall(r"[^\x00-\x7F\u2022\u2013\u2014\u00A0]", text))
    diag.non_standard_char_count = non_standard
    if non_standard > 25:
        diag.formatting_warnings.append(f"Detected {non_standard} non-standard characters or symbols.")

    return text, diag

