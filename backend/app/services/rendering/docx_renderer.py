"""Deterministic ATS single-column DOCX generator using python-docx."""

from __future__ import annotations

import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.schemas.rendering import DeterministicResumeDocument


class ATSDocxRenderer:
    """Renders a structured, validated DeterministicResumeDocument into an ATS-safe single-column DOCX file."""

    def __init__(self, doc: DeterministicResumeDocument) -> None:
        self.doc = doc
        self.config = doc.config

    def render_docx_bytes(self) -> bytes:
        """Produces clean, semantic DOCX binary bytes with standard styles and margins."""
        document = Document()

        # Set page margins
        sections = document.sections
        for section in sections:
            section.top_margin = Inches(self.config.margins_inch)
            section.bottom_margin = Inches(self.config.margins_inch)
            section.left_margin = Inches(self.config.margins_inch)
            section.right_margin = Inches(self.config.margins_inch)

        # 1. Candidate Name Header
        name_para = document.add_paragraph()
        name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_para.paragraph_format.space_before = Pt(0)
        name_para.paragraph_format.space_after = Pt(2)
        name_run = name_para.add_run(self.doc.header.name)
        name_run.font.name = self.config.font_family
        name_run.font.size = Pt(18)
        name_run.font.bold = True
        name_run.font.color.rgb = RGBColor(17, 24, 39)

        # 2. Contact Information Line
        contact_parts: list[str] = [self.doc.header.email]
        if self.doc.header.phone:
            contact_parts.append(self.doc.header.phone)
        if self.doc.header.location:
            contact_parts.append(self.doc.header.location)
        if self.doc.header.linkedin:
            contact_parts.append(self.doc.header.linkedin)
        if self.doc.header.github:
            contact_parts.append(self.doc.header.github)
        if self.doc.header.portfolio:
            contact_parts.append(self.doc.header.portfolio)

        contact_para = document.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_para.paragraph_format.space_before = Pt(0)
        contact_para.paragraph_format.space_after = Pt(10)
        contact_run = contact_para.add_run(" | ".join(contact_parts))
        contact_run.font.name = self.config.font_family
        contact_run.font.size = Pt(9.5)
        contact_run.font.color.rgb = RGBColor(75, 85, 99)

        # 3. Section Blocks
        for sec in self.doc.sections:
            if not sec.items:
                continue

            # Section Heading
            heading_para = document.add_paragraph()
            heading_para.paragraph_format.space_before = Pt(8)
            heading_para.paragraph_format.space_after = Pt(3)
            heading_para.paragraph_format.keep_with_next = True
            h_run = heading_para.add_run(sec.heading.upper())
            h_run.font.name = self.config.font_family
            h_run.font.size = Pt(self.config.heading_size_pt)
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(31, 41, 55)

            # Items
            for item in sec.items:
                if item.title:
                    title_para = document.add_paragraph()
                    title_para.paragraph_format.space_before = Pt(4)
                    title_para.paragraph_format.space_after = Pt(2)
                    title_para.paragraph_format.keep_with_next = True
                    t_run = title_para.add_run(item.title)
                    t_run.font.name = self.config.font_family
                    t_run.font.size = Pt(self.config.font_size_pt + 0.5)
                    t_run.font.bold = True
                    t_run.font.color.rgb = RGBColor(17, 24, 39)

                if item.raw_text:
                    body_para = document.add_paragraph()
                    body_para.paragraph_format.space_before = Pt(0)
                    body_para.paragraph_format.space_after = Pt(4)
                    b_run = body_para.add_run(item.raw_text)
                    b_run.font.name = self.config.font_family
                    b_run.font.size = Pt(self.config.font_size_pt)
                    b_run.font.color.rgb = RGBColor(55, 65, 81)

                for bullet in item.bullets:
                    bullet_para = document.add_paragraph(style="List Bullet")
                    bullet_para.paragraph_format.space_before = Pt(0)
                    bullet_para.paragraph_format.space_after = Pt(2)
                    bullet_run = bullet_para.add_run(bullet)
                    bullet_run.font.name = self.config.font_family
                    bullet_run.font.size = Pt(self.config.font_size_pt)
                    bullet_run.font.color.rgb = RGBColor(55, 65, 81)

        buffer = io.BytesIO()
        document.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()
        return docx_bytes
